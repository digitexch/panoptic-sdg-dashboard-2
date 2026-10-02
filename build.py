"""Build the Panoptic SDG dashboard from the department workbook.

    Workbook (edit here)  ->  build.py  ->  panoptic_dashboard_data.json  ->  panoptic-sdg-dashboard.html

Usage:  python3 build.py [workbook.xlsx]
The workbook is the only place data is edited. This script reads it, checks it,
writes the JSON feed and injects that feed into dashboard.template.html.
"""
import json, sys, os, datetime
from collections import OrderedDict
from openpyxl import load_workbook
import warnings; warnings.filterwarnings("ignore", module="openpyxl")   # Excel extensions (e.g. data bars) are read-only noise here

SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join("data", "Panoptic_SDG_database_department.xlsx")
# Public build (set by the GitHub workflow): evaluator names, comments and the project lead are left out of the page
PUBLIC = os.environ.get("PANOPTIC_PUBLIC") == "1"
# Release switches: what the dashboard shows. The workbook keeps all data either way.
FEATURES = {
    "indicators": False,   # SDG indicators: held back for a later release
    "scoring": True,       # project scoring layer
}
TEMPLATE, OUT_JSON, OUT_HTML = "dashboard.template.html", "panoptic_dashboard_data.json", "panoptic-sdg-dashboard.html"

wb = load_workbook(SRC)  # raw values: the build never depends on cached formula results
def rows(sheet, first_col_required=True):
    ws = wb[sheet]; head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if first_col_required and (r[0] is None or str(r[0]).strip() == ""): continue
        yield dict(zip(head, r))
txt = lambda v: "" if v is None else str(v).strip()
errors = []
def check(cond, msg):
    if not cond: errors.append(msg)

# 1. Principles and sub-principles
principles = [r for r in rows("Principles") if txt(r["Principle"]).startswith("PP")]
subs = [r for r in rows("Sub-principles") if txt(r["Sub-principle"]).startswith("PP")]
pids = {txt(p["Principle"]) for p in principles}
sub_ids = [txt(s["Sub-principle"]) for s in subs]
check(len(sub_ids) == len(set(sub_ids)), "Duplicate sub-principle IDs")
for s in subs:
    check(txt(s["Principle"]) in pids, f"{s['Sub-principle']}: unknown principle {s['Principle']}")
    check(isinstance(s["Weight"], (int, float)) and s["Weight"] > 0, f"{s['Sub-principle']}: weight missing")
total_w = sum(s["Weight"] or 0 for s in subs)
check(abs(total_w - 100) < 1e-9, f"Sub-principle weights add up to {total_w}, not 100")

# 2. SDG goals
goals = [r for r in rows("SDG goals") if isinstance(r["SDG"], (int, float))]
check(len(goals) == 17, f"Expected 17 SDGs, found {len(goals)}")

# 3. Target links (the core table) and indicator links
links, targets, seen = [], OrderedDict(), set()
for r in rows("Target links"):
    lid, sub, t, a = txt(r["Link ID"]), txt(r["Sub-principle"]), txt(r["Target"]), txt(r["Alignment"])
    check(sub in sub_ids, f"{lid}: sub-principle {sub} not on the Sub-principles sheet")
    check(a in ("Primary", "Secondary"), f"{lid}: alignment must be Primary or Secondary, found '{a}'")
    check(lid == f"{sub}|{t}", f"{lid}: Link ID should be {sub}|{t}")
    check(lid not in seen, f"{lid}: duplicate link"); seen.add(lid)
    g = int(r["SDG"]); check(t.split(".")[0] == str(g), f"{lid}: target {t} is not under SDG {g}")
    check(txt(r["Target text (UNSD)"]) != "", f"{lid}: target text missing")
    targets.setdefault(t, {"g": g, "s": txt(r["Target short name"]), "t": txt(r["Target text (UNSD)"])})
    links.append({"s": sub, "t": t, "a": "P" if a == "Primary" else "S", "src": "E", "l": txt(r["Department label"]), "r": ""})
inds, also = {}, {}
for r in rows("Indicator links"):
    t, ic = txt(r["Target"]), txt(r["Indicator"])
    lst = inds.setdefault(t, [])
    if ic not in [x[0] for x in lst]: lst.append([ic, txt(r["Indicator text (UNSD)"])])
    if txt(r["Also listed under"]): also[ic] = txt(r["Also listed under"])
for t in targets: check(t in inds and inds[t], f"Target {t}: no indicators on the Indicator links sheet")

# 4. Projects and evaluator scores
projects, seen_p = [], {}
for r in rows("Projects"):
    pid = txt(r["Project ID"])
    check(pid not in seen_p, f"Project {pid}: duplicate ID")
    check(txt(r["Project name"]) != "", f"Project {pid}: name missing")
    d = r["Assessment date"]
    P = {"id": pid, "name": txt(r["Project name"]), "dept": txt(r["Department"]), "loc": txt(r["Location"]), "sector": txt(r["Sector"]),
         "stage": txt(r["Stage gate"]), "lead": txt(r["Project lead"]), "date": d.strftime("%-d %B %Y") if hasattr(d, "strftime") else txt(d), "evals": []}
    seen_p[pid] = P; projects.append(P)
seen_s = set()
for r in rows("Project scores"):
    pid, ev, sub, sc = txt(r["Project ID"]), txt(r["Evaluator"]), txt(r["Sub-principle"]), r["Score (1–5)"]
    where = f"Project scores: {pid} / {ev or '?'} / {sub or '?'}"
    if not (ev or sub or sc is not None): continue
    check(pid in seen_p, f"{where}: project ID not on the Projects sheet")
    check(ev != "", f"{where}: evaluator missing")
    check(sub in sub_ids, f"{where}: unknown sub-principle")
    ok = isinstance(sc, (int, float)) and 1 <= sc <= 5 and abs(sc * 2 - round(sc * 2)) < 1e-9
    check(ok, f"{where}: score must be 1 to 5 in steps of 0.5, found {sc!r}")
    k = (pid, ev, sub); check(k not in seen_s, f"{where}: scored twice"); seen_s.add(k)
    if pid not in seen_p or not ok or sub not in sub_ids: continue
    P = seen_p[pid]; E = next((e for e in P["evals"] if e["n"] == ev), None)
    if not E: E = {"n": ev, "s": {}, "c": {}}; P["evals"].append(E)
    E["s"][sub] = sc
    if txt(r["Comment"]): E["c"][sub] = txt(r["Comment"])

if errors:
    print("BUILD STOPPED: fix these in the workbook, then run again:"); [print("  -", e) for e in errors]; sys.exit(1)

# 5. Feed in the shape the dashboard reads
pw = {}
for s in subs: pw[txt(s["Principle"])] = pw.get(txt(s["Principle"]), 0) + s["Weight"]
data = {
    "principles": [{"id": txt(p["Principle"]), "name": txt(p["Name"]), "short": txt(p["Short name"]), "w": pw.get(txt(p["Principle"]), 0)} for p in principles],
    "subs": [{"id": txt(s["Sub-principle"]), "p": txt(s["Principle"]), "name": txt(s["Name"]), "short": txt(s["Short name"]), "w": s["Weight"],
              "crit": txt(s["Key policy criterion"]), "q": txt(s["Assessment question"]), "why": txt(s["Alignment rationale"]), "pres": txt(s["Prescripts"])} for s in subs],
    "sdgs": [{"n": int(g["SDG"]), "name": txt(g["Short name"]), "title": txt(g["Goal title"]), "c": txt(g["Colour"])} for g in goals],
    "targets": dict(targets), "links": links,
    "meta": {"framework": "UNSD global indicator framework, after the March 2026 refinement",
             "source": "department Panoptic-to-SDG alignment (primary and secondary columns)",
             "prepared": datetime.date.today().strftime("%-d %B %Y"), "workbook": SRC, "features": FEATURES},
}
if FEATURES["indicators"]: data.update(inds=inds, also=also)
if FEATURES["scoring"]:
    if PUBLIC:
        for P in projects:
            P["lead"] = ""
            for i, E in enumerate(P["evals"], 1): E["n"], E["c"] = f"Evaluator {i}", {}
    data["projects"] = projects
data["meta"]["public"] = PUBLIC
with open(OUT_JSON, "w", encoding="utf-8") as f: json.dump(data, f, ensure_ascii=False, indent=1)
payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
html = open(TEMPLATE, encoding="utf-8").read()
assert html.count("__DATA__") == 1, "Template must contain exactly one __DATA__ placeholder"
# Visual assets (kept out of the data feed): SDG icons from assets/sdg/sdg-NN.png
import base64, os
icons = {}
for n in range(1, 18):
    f = os.path.join("assets", "sdg", f"sdg-{n:02d}.png")
    if os.path.exists(f): icons[n] = "data:image/png;base64," + base64.b64encode(open(f, "rb").read()).decode()
assets = json.dumps({"sdg": icons}, separators=(",", ":"))
assert html.count("__ASSETS__") == 1, "Template must contain exactly one __ASSETS__ placeholder"
open(OUT_HTML, "w", encoding="utf-8").write(html.replace("__DATA__", payload).replace("__ASSETS__", assets))
np = sum(l["a"] == "P" for l in links)
print(f"OK  {len(principles)} principles, {len(subs)} sub-principles (weights {total_w:g}), {len(links)} links ({np} primary, {len(links)-np} secondary), "
      f"{len(targets)} targets in {len({t['g'] for t in targets.values()})} SDGs, {sum(map(len, inds.values()))} indicators"
      f" ({'shown' if FEATURES['indicators'] else 'held back'}), {len(projects)} projects / {len(seen_s)} scores")
print(f"    wrote {OUT_JSON} and {OUT_HTML} ({len(icons)}/17 SDG icons)")
# Documentation diagram, regenerated from the same feed so its numbers stay current
if os.path.exists(os.path.join("docs", "build_framework.py")):
    import runpy; runpy.run_path(os.path.join("docs", "build_framework.py"))
